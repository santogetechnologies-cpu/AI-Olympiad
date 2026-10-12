import os
import docx

senior_dir = "/Users/ashlinjeshma/Desktop/Project/public/Senior Level "

c11_files = [
    "AI_Basics_Grade11_Month1_Workbook.docx",
    "Grade11_AI_Olympiad_Month2_Unit2_Communicate_with_AI_Workbook.docx",
    "AI_Olympiad_Grade11_Month3_Unit3_AI_Applications_Workbook.docx",
    "Class11_AI_Olympiad_Unit4_AI_and_Career_Workbook.docx",
    "Class11_AI_Olympiad_Unit5_AI_Tools_Workbook.docx",
    "Class11_AI_Olympiad_Unit6_Responsible_AI_Workbook.docx",
]

c12_files = [
    "World_of_Generative_AI_Understanding_LLMs (3).docx",
    "Prompt_Plan_Produce_AI_Workflow_Basics_Class12.docx",
    "AI_for_Innovation_AI_for_Enterprise_Class12.docx",
    "Your_AI_Career_Launchpad_Build_Your_Professional_Profile_Class12.docx",
    "AI_Assistant_Lab_AI_Project_Studio_Class12.docx",
    "Verify_Before_You_Trust.docx",
]

def analyze_docx_file(fname):
    path = os.path.join(senior_dir, fname)
    doc = docx.Document(path)
    print("=" * 80)
    print(f"ANALYZING: {fname}")
    print("=" * 80)
    # Collect all headings or bold titles or numbered sections
    major_sections = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        is_heading = False
        try:
            if p.style and p.style.name and p.style.name.startswith("Heading"):
                is_heading = True
        except:
            pass
        if any(run.bold for run in p.runs[:2]) and len(t) < 80:
            is_heading = True
        if any(t.startswith(prefix) for prefix in ["Chapter", "Lesson", "Unit", "Part", "Topic", "1.", "2.", "3.", "4.", "5.", "6."]) and len(t) < 80:
            is_heading = True
        if is_heading:
            major_sections.append(t)
    print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")
    print(f"Major sections ({len(major_sections)}):")
    for s in major_sections[:20]:
        print(f"  - {s}")
    print()

print("--- CLASS 11 WORKBOOKS ---")
for f in c11_files:
    analyze_docx_file(f)

print("--- CLASS 12 FILES ---")
for f in c12_files:
    analyze_docx_file(f)
