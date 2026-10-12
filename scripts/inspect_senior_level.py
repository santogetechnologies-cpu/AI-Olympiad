import os
import pymupdf
import docx

senior_dir = "/Users/ashlinjeshma/Desktop/Project/public/Senior Level "

files = sorted(os.listdir(senior_dir))

print(f"Total files in '{senior_dir}': {len(files)}\n")

for f in files:
    full_path = os.path.join(senior_dir, f)
    if os.path.isdir(full_path):
        continue
    
    print("=" * 80)
    print(f"FILE: {f} (size: {os.path.getsize(full_path)} bytes)")
    print("=" * 80)
    
    if f.endswith(".pdf"):
        doc = pymupdf.open(full_path)
        print(f"Type: PDF, Pages: {len(doc)}")
        for i in range(min(5, len(doc))):
            text = doc[i].get_text().strip()
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            header = " | ".join(lines[:3]) if lines else "EMPTY"
            print(f"  Page {i+1}: {header}")
        # print table of contents or major headings
        all_text = ""
        for page in doc:
            all_text += page.get_text() + "\n"
        print(f"  Total characters: {len(all_text)}")
        headings = [line.strip() for line in all_text.splitlines() if any(k in line.upper() for k in ["CHAPTER", "LESSON", "UNIT", "MODULE", "TABLE OF CONTENTS", "CONTENT", "TOPIC"]) and len(line.strip()) < 80]
        print(f"  Key headings sampled: {headings[:15]}")
        
    elif f.endswith(".docx"):
        try:
            doc = docx.Document(full_path)
            print(f"Type: DOCX, Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
            headings = []
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt and (p.style.name.startswith("Heading") or any(k in txt.upper() for k in ["CHAPTER", "LESSON", "UNIT", "MODULE", "MONTH", "GRADE", "CLASS", "OBJECTIVE", "TOPIC"])):
                    if len(txt) < 100:
                        headings.append(txt)
            print(f"  Sample Headings ({len(headings)} found): {headings[:12]}")
        except Exception as e:
            print(f"  Error reading docx: {e}")
    else:
        print(f"  Other file type: {f}")
    print()
