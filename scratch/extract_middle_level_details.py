import os
import docx
from pypdf import PdfReader

dir_path = r"c:\Users\aswin\Desktop\AI-Olympiad\public\Middle Level"
out_path = r"c:\Users\aswin\Desktop\AI-Olympiad\scratch\extracted_middle_level_summary.txt"

with open(out_path, "w", encoding="utf-8") as out:
    for fname in sorted(os.listdir(dir_path)):
        if fname.endswith(".jpeg"):
            continue
        fpath = os.path.join(dir_path, fname)
        out.write(f"\n=======================================================\n")
        out.write(f"FILE: {fname}\n")
        out.write(f"=======================================================\n")
        
        if fname.endswith(".docx"):
            try:
                doc = docx.Document(fpath)
                out.write(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}\n")
                out.write("--- PARAGRAPHS ---\n")
                for i, p in enumerate(doc.paragraphs):
                    txt = p.text.strip()
                    if txt:
                        # check if heading or bold
                        is_bold = any(r.bold for r in p.runs if r.bold is not None)
                        style = p.style.name if p.style else ""
                        out.write(f"P[{i+1}][{style}]{'[BOLD] ' if is_bold else ''}{txt}\n")
                
                out.write("--- SAMPLE TABLES ---\n")
                for t_idx, tbl in enumerate(doc.tables[:10]):
                    out.write(f"\nTable {t_idx+1} ({len(tbl.rows)} rows, {len(tbl.columns)} cols):\n")
                    for r in tbl.rows[:5]:
                        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
                        out.write(f"  | {' | '.join(row_txt)} |\n")
            except Exception as e:
                out.write(f"Error reading docx: {e}\n")
                
        elif fname.endswith(".pdf"):
            try:
                reader = PdfReader(fpath)
                out.write(f"Total pages: {len(reader.pages)}\n")
                for p_idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    out.write(f"\n--- PAGE {p_idx+1} ---\n")
                    for line in text.split("\n"):
                        if line.strip():
                            out.write(f"{line.strip()}\n")
            except Exception as e:
                out.write(f"Error reading pdf: {e}\n")

print(f"Extraction complete! Written to {out_path}")
