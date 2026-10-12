import os
import docx
from pypdf import PdfReader

dir_path = r"c:\Users\aswin\Desktop\AI-Olympiad\public\Middle Level"

print("--- Inspecting Middle Level Files ---")
for fname in sorted(os.listdir(dir_path)):
    fpath = os.path.join(dir_path, fname)
    print(f"\n==========================================")
    print(f"FILE: {fname} (Size: {os.path.getsize(fpath)} bytes)")
    print(f"==========================================")
    
    if fname.endswith(".docx"):
        try:
            doc = docx.Document(fpath)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            print(f"Total paragraphs: {len(paragraphs)}")
            print("First 20 non-empty paragraphs:")
            for i, p in enumerate(paragraphs[:20]):
                print(f"  [{i+1}] {p[:120]}")
            # Also check headings / tables
            print(f"Total tables: {len(doc.tables)}")
            headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading") or (p.runs and p.runs[0].bold and len(p.text.strip()) < 80)]
            print(f"Detected headings/bold items (first 15):")
            for h in headings[:15]:
                print(f"  * {h}")
        except Exception as e:
            print(f"Error reading docx: {e}")
            
    elif fname.endswith(".pdf"):
        try:
            reader = PdfReader(fpath)
            print(f"Total pages: {len(reader.pages)}")
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                lines = [l.strip() for l in text.split("\n") if l.strip()]
                print(f"--- Page {idx+1} ({len(lines)} lines) ---")
                for l in lines[:10]:
                    print(f"  {l[:100]}")
        except Exception as e:
            print(f"Error reading pdf: {e}")
